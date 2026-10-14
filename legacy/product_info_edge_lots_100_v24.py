# -*- coding: utf-8 -*-
"""Edge Product Info: lots manuels de 100; journal durable; aucun ancien journal repris."""
from __future__ import annotations
import argparse, csv, hashlib, json, os, re, sqlite3, sys, time, uuid
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

ROOT=Path(__file__).resolve().parent
BOOK=ROOT/'Benchmark_Product_Info_1500_questions_revision_naturelle.xlsx'
PROFILE=ROOT/'.edge_product_info_profile_lots100_v24'
LATEST=ROOT/'product_info_edge_lots100_v24_v21_latest.json'
BATCH=100
INPUTS=('[contenteditable="true"][aria-label="Message Product Info Assistant"]',
        '[contenteditable="true"][data-placeholder="Message Product Info Assistant"]',
        '[role="textbox"][aria-label="Message Product Info Assistant"]',
        '[contenteditable="true"][aria-label*="Message Copilot" i]',
        '[contenteditable="true"][data-placeholder*="Message Copilot" i]',
        '[role="textbox"][aria-label*="Message Copilot" i]',
        'textarea[placeholder*="Message Copilot" i]',
        'textarea[aria-label*="message" i]')
ANSWERS=('[data-message-author-role="assistant"]','[data-author="assistant"]',
         '[data-testid*="assistant-message" i]','[aria-label*="Product Info Assistant said" i]',
         '[role="article"][aria-label*="assistant" i]','[role="article"]')
WAIT=('mise en place en cours','lining things up','nous travaillons sur une réponse',
      'veuillez patienter','recherche en cours','génération en cours','generation en cours')
STOP=re.compile('arrêter|interrompre|stop generating|stop responding',re.I)

def stamp(): return datetime.now(timezone.utc).isoformat()
def norm(s): return re.sub(r'\s+',' ',str(s or '').replace('\xa0',' ')).strip()
def waiting(s):
    t=norm(s).casefold()
    return any(x in t for x in WAIT) or t in ('product info assistant','product info assistant said:')

def catalogue():
    if not BOOK.is_file(): raise FileNotFoundError(f'Classeur absent : {BOOK.name}')
    wb=load_workbook(BOOK,read_only=True,data_only=True)
    try:
        if 'QUESTIONS_1500' not in wb: raise ValueError('Onglet QUESTIONS_1500 absent')
        it=wb['QUESTIONS_1500'].iter_rows(values_only=True); head=list(next(it))
        keys=['ID question','ID conversation','Tour','Question utilisateur']
        if any(head.count(k)!=1 for k in keys): raise ValueError('En-têtes manquants ou dupliqués')
        rows=[]
        for values in it:
            r=dict(zip(head,values)); q=str(r.get('Question utilisateur') or '').strip()
            if q:
                rows.append(dict(qid=norm(r['ID question']),cid=norm(r['ID conversation']),
                    tour=int(r['Tour']),question=q,produit=norm(r.get('Produit cible exact')),
                    champ=norm(r.get('Champ exact')),attendu=norm(r.get('Valeur source brute')),
                    gt=norm(r.get('Statut GT'))))
    finally: wb.close()
    if not rows or [r['qid'] for r in rows]!=[f'Q{i:04d}' for i in range(1,len(rows)+1)]:
        raise ValueError('IDs manquants, dupliqués ou hors ordre')
    digest=hashlib.sha256(json.dumps(rows,sort_keys=True,ensure_ascii=False).encode()).hexdigest()
    return rows,digest

def open_db(path):
    db=sqlite3.connect(path,timeout=30)
    db.execute('PRAGMA journal_mode=WAL'); db.execute('PRAGMA synchronous=FULL')
    db.execute('CREATE TABLE IF NOT EXISTS meta(k TEXT PRIMARY KEY,v TEXT)')
    db.execute('CREATE TABLE IF NOT EXISTS turns(qid TEXT PRIMARY KEY,cid TEXT,tour INTEGER,question TEXT,produit TEXT,champ TEXT,attendu TEXT,gt TEXT,status TEXT,answer TEXT,started TEXT,finished TEXT,latency REAL,error TEXT,batch INTEGER)')
    db.commit(); return db

COLS='qid,cid,tour,question,produit,champ,attendu,gt,status,answer,started,finished,latency,error,batch'
def csv_export(db,path):
    dest=path.with_suffix('.csv'); temp=dest.with_suffix('.csv.tmp')
    with temp.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f); w.writerow(COLS.split(',')); w.writerows(db.execute('SELECT '+COLS+' FROM turns ORDER BY qid'))
        f.flush(); os.fsync(f.fileno())
    os.replace(temp,dest)

def record(db,r,status,started,answer=None,error=None,latency=None):
    with db:
        db.execute('INSERT INTO turns VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(qid) DO UPDATE SET status=excluded.status,answer=excluded.answer,finished=excluded.finished,latency=excluded.latency,error=excluded.error',
            (r['qid'],r['cid'],r['tour'],r['question'],r['produit'],r['champ'],r['attendu'],r['gt'],
             status,answer,started,stamp(),latency,error,(int(r['qid'][1:])-1)//BATCH+1))

def composer(page):
    for css in INPUTS:
        hits=[]
        for frame in page.frames:
            try:
                loc=frame.locator(css)
                hits.extend(loc.nth(i) for i in range(loc.count()) if loc.nth(i).is_visible())
            except Exception: pass
        if hits:
            if len(hits)!=1: raise RuntimeError('Plusieurs zones de saisie; fermer les onglets supplémentaires')
            return hits[0]
    raise RuntimeError('Champ Message Product Info Assistant absent')

def editor_value(box):
    try:
        value=box.input_value(timeout=1200)
        if norm(value): return norm(value)
    except Exception: pass
    try: return norm(box.inner_text(timeout=1200))
    except Exception: return ''

def page_chat(ctx, wait_seconds=45):
    end=time.monotonic()+wait_seconds
    last_count=0
    while time.monotonic()<end:
        pages=[]
        for p in ctx.pages:
            if not p.is_closed():
                try: composer(p); pages.append(p)
                except RuntimeError: pass
        last_count=len(pages)
        if len(pages)==1: return pages[0]
        if len(pages)>1: raise RuntimeError(f'{len(pages)} onglets avec zone de saisie; fermer les onglets supplémentaires')
        time.sleep(.8)
    titles=[]
    for p in ctx.pages:
        if not p.is_closed():
            try: titles.append(p.title()[:80])
            except Exception: titles.append('[titre inaccessible]')
    raise RuntimeError(f'0 onglet avec zone Message Copilot/Product Info après attente. Onglets pilotés: {titles}')

def answers(page,questions):
    choices=[]
    for frame in page.frames:
        for css in ANSWERS:
            try:
                loc=frame.locator(css); vals=[]
                for i in range(loc.count()):
                    node=loc.nth(i)
                    if not node.is_visible(): continue
                    text=node.inner_text(timeout=1200); label=node.get_attribute('aria-label') or ''
                    if css=='[role="article"]' and 'assistant' not in label.casefold() and 'product info assistant' not in text.casefold(): continue
                    t=text.strip()
                    if norm(t) and norm(t) not in questions: vals.append(t)
                if vals:
                    choices.append((css,vals)); break
            except Exception: continue
    if len(choices)>1: raise RuntimeError('Réponses détectées dans plusieurs frames; attribution incertaine')
    return choices[0] if choices else (None,[])

def sent(page,q):
    for frame in page.frames:
        try:
            loc=frame.get_by_text(q,exact=True)
            for i in range(loc.count()):
                el=loc.nth(i)
                if el.is_visible() and not el.evaluate('(e)=>!!e.closest("[contenteditable=true]")'):
                    return True
        except Exception: pass
    return False

def generating(page):
    for f in page.frames:
        try:
            loc=f.get_by_role('button',name=STOP)
            if any(loc.nth(i).is_visible() for i in range(loc.count())): return True
        except Exception: pass
    return False

def draft_value(box):
    for read in (
        lambda: box.input_value(timeout=1000),
        lambda: box.inner_text(timeout=1000),
        lambda: box.evaluate('(e)=>e.innerText || e.textContent || ""'),
    ):
        try:
            value=norm(read())
            if value and value.casefold() not in ('message product info assistant','message copilot'):
                return value
        except Exception: pass
    return ''

def send_buttons(page):
    pattern=re.compile(r'envoyer|send|submit|soumettre',re.I)
    found=[]; seen=set()
    for frame in page.frames:
        for css in ('button[aria-label]','button[title]','[role="button"][aria-label]'):
            try:
                loc=frame.locator(css)
                for i in range(loc.count()):
                    b=loc.nth(i)
                    if not b.is_visible() or not b.is_enabled(): continue
                    label=' '.join(x or '' for x in (b.get_attribute('aria-label'),b.get_attribute('title')))
                    if not pattern.search(label): continue
                    bb=b.bounding_box(); key=str(bb)
                    if key not in seen: seen.add(key); found.append(b)
            except Exception: pass
    return found

def dispatch_signal(page,q,before):
    if sent(page,q): return True
    try:
        if not draft_value(composer(page)):
            css,vals=answers(page,set())
            if len(vals)>len(before) or (vals and before and norm(vals[-1])!=norm(before[-1])):
                return True
            # Le composeur vide juste apres Enter est aussi un signal d'envoi sur cette interface.
            return True
    except Exception: pass
    return False

def wait_dispatch(page,q,before,seconds):
    end=time.monotonic()+seconds
    while time.monotonic()<end:
        if dispatch_signal(page,q,before): return True
        time.sleep(.35)
    return False

def send_once(page,q,before):
    """Saisie puis au maximum deux gestes: Enter, et un clic uniquement si le brouillon exact reste present."""
    box=composer(page)
    if draft_value(box): raise RuntimeError('Brouillon present; aucun ecrasement')
    box.click(); box.press('Control+A'); box.press('Backspace')
    box.press_sequentially(q,delay=1)
    typed=draft_value(box)
    if norm(typed)!=norm(q):
        raise RuntimeError(f'Saisie non confirmee ({len(typed)}/{len(q)} caracteres); aucun envoi')
    box.press('Enter')
    if wait_dispatch(page,q,before,12): return
    # Enter peut inserer une ligne ou rester sans effet. Un clic est autorise seulement si le brouillon exact est intact.
    box=composer(page); typed=draft_value(box)
    if norm(typed)!=norm(q):
        raise RuntimeError('Etat du composeur ambigu apres Entree; aucun second geste')
    buttons=send_buttons(page)
    if len(buttons)!=1:
        raise RuntimeError(f'Bouton Envoyer non unique ({len(buttons)}); aucun clic')
    buttons[0].click()
    if wait_dispatch(page,q,before,20): return
    raise RuntimeError('Envoi non confirme apres Enter puis clic unique')

def wait_answer(page,before,q,questions,timeout,stable):
    end=time.monotonic()+timeout
    prev=before[-1] if before else ''
    last=''; since=0; chosen=None
    while time.monotonic()<end:
        css,vals=answers(page,questions)
        if vals:
            text=vals[-1]
            fresh=(len(vals)>len(before) or bool(prev and norm(text)!=norm(prev)) or not before)
            if fresh and not waiting(text):
                if text!=last:
                    last=text; since=time.monotonic(); chosen=css
                elif time.monotonic()-since>=stable and not generating(page):
                    css2,vals2=answers(page,questions)
                    if vals2 and vals2[-1]==text and not waiting(text) and not generating(page):
                        return text,chosen
            else:
                last=''; since=0
        time.sleep(.8)
    raise TimeoutError('Reponse finale non confirmee; aucune reponse attribuee')

def user_messages(page):
    for frame in page.frames:
        for css in ('[data-message-author-role="user"]','[data-author="user"]','[class*="fai-UserMessage"]'):
            try:
                loc=frame.locator(css)
                found=[loc.nth(i) for i in range(loc.count()) if loc.nth(i).is_visible() and norm(loc.nth(i).inner_text(timeout=1200))]
                if found: return len(found)
            except Exception: pass
    return 0

def new_chat(ctx,first,questions):
    if first: print('Dans CETTE fenetre Edge, cliquez sur Product Info Assistant puis ouvrez un chat VIERGE.')
    else: print('Lot de 100 sauvegardé. Ouvrez VOUS-MÊME un NOUVEAU chat vierge Product Info dans CETTE fenêtre Edge.')
    input('Quand la zone Message Product Info Assistant est visible, appuyez une fois sur Entree : ')
    print('Détection du chat en cours...', flush=True)
    p=page_chat(ctx, wait_seconds=60)
    if answers(p,questions)[1] or user_messages(p) or draft_value(composer(p)):
        raise RuntimeError('Chat non vierge ou brouillon présent; aucun envoi')
    return p

def run(a):
    rows,digest=catalogue(); questions={r['question'] for r in rows}
    print(f'Catalogue: {len(rows)} questions; pause après chaque lot de {BATCH}.')
    if a.dry_run:
        print('Première:',rows[0]['qid'],rows[0]['question']); print('Dernière:',rows[-1]['qid'],'; aucun envoi'); return
    if a.resume:
        if not LATEST.exists(): raise RuntimeError('Aucune campagne à reprendre')
        path=ROOT/json.loads(LATEST.read_text(encoding='utf-8'))['db']
        if not path.exists(): raise RuntimeError('Base de campagne absente')
    else:
        path=ROOT/('product_info_edge_lots100_v24_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'_'+uuid.uuid4().hex[:6]+'.sqlite3')
    db=open_db(path)
    try:
        old=db.execute("SELECT v FROM meta WHERE k='catalogue'").fetchone()
        if old and old[0]!=digest: raise RuntimeError('Classeur modifié depuis le début de cette campagne')
        if not old:
            with db: db.execute("INSERT INTO meta VALUES('catalogue',?)",(digest,))
        states=list(db.execute('SELECT qid,status FROM turns ORDER BY qid'))
        if any(s!='CAPTURED' for _,s in states): raise RuntimeError('Tour incertain; ne pas reprendre automatiquement')
        done=len(states)
        if [q for q,_ in states]!=[r['qid'] for r in rows[:done]]: raise RuntimeError('Journal non contigu')
        end=len(rows) if a.all else min(done+a.pilot,len(rows))
        if done>=end: print('Aucune question restante'); return
        if not a.resume: LATEST.write_text(json.dumps({'db':path.name}),encoding='utf-8')
        print('Base:',path.name,'capturées:',done,'à traiter:',end-done)
        with sync_playwright() as pw:
            ctx=pw.chromium.launch_persistent_context(str(PROFILE),channel='msedge',headless=False,viewport=None,args=['--start-maximized'])
            try:
                page=None
                for i in range(done,end):
                    r=rows[i]
                    if i%BATCH==0: page=new_chat(ctx,i==0,questions)
                    elif page is None:
                        input('Rouvrez dans CET Edge le chat du dernier tour capturé, puis Entrée : ')
                        page=page_chat(ctx)
                        if not sent(page,rows[i-1]['question']): raise RuntimeError('Dernière question non visible; mauvais chat ou historique masqué')
                    before=answers(page,questions)[1]; started=stamp(); start=time.monotonic()
                    record(db,r,'PENDING',started); csv_export(db,path)
                    try:
                        send_once(page,r['question'],before)
                        text,css=wait_answer(page,before,r['question'],questions,a.timeout,a.stable)
                        record(db,r,'CAPTURED',started,answer=text,latency=round(time.monotonic()-start,2))
                        csv_export(db,path)
                        print(f'{r["qid"]} CAPTURED {i+1}/{end} ({len(text)} caractères, {css})',flush=True)
                    except BaseException as e:
                        record(db,r,'UNCERTAIN',started,error=f'{type(e).__name__}: {e}'); csv_export(db,path)
                        try: page.screenshot(path=str(path.with_name(path.stem+'_'+r['qid']+'_erreur.png')))
                        except Exception: pass
                        print(f'ARRÊT {r["qid"]} INCERTAIN: {type(e).__name__}: {e}',flush=True)
                        raise
                print('Terminé. Résultats:',path.name,'et',path.with_suffix('.csv').name)
            finally:
                try: ctx.close()
                except Exception: print('Fermeture Edge incomplète; base et CSV conservés.')
    finally: db.close()

def main():
    ap=argparse.ArgumentParser(description='Edge: 100 questions par chat, pause manuelle entre lots')
    mode=ap.add_mutually_exclusive_group(required=True)
    mode.add_argument('--dry-run',action='store_true'); mode.add_argument('--run',action='store_true')
    ap.add_argument('--all',action='store_true'); ap.add_argument('--pilot',type=int,default=5)
    ap.add_argument('--resume',action='store_true',help='reprendre une campagne sans tour incertain')
    ap.add_argument('--timeout',type=int,default=360); ap.add_argument('--stable',type=float,default=8)
    a=ap.parse_args()
    if a.pilot<1 or a.timeout<30 or a.stable<3: ap.error('Paramètres invalides')
    try: run(a)
    except KeyboardInterrupt: print('Interrompu; réponses déjà capturées conservées.'); sys.exit(130)
    except Exception as e: print('ERREUR:',e); sys.exit(1)
if __name__=='__main__': main()
