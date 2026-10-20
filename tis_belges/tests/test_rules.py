"""Tests des règles métier. Lancer : python -m unittest discover -s tests -v

Le test sur fichiers réels n'est exécuté que si TIS_SAMPLE_DIR pointe vers le dossier
contenant les exemples (TISBE_PROD, AM_File, historique Valor)."""

import os
import unittest
from datetime import date
from pathlib import Path

import openpyxl

import rules

RULES = {
    "min_rows": 3,
    "max_rows": 100,
    "currencies": ["EUR", "USD", "CHF", "GBP"],
    "types": ["TIBR", "TISF"],
    "allow_duplicates": False,
}
TARGET = "20261016"


def line(isin="FR0010523191", name="R-co Conviction Club", ccy="EUR", day=TARGET,
         nav=142.84, tis=1.32, typ="TIBR"):
    return (isin, name, ccy, day, None, nav, tis, typ)


def good_rows():
    return [
        line("FR0010523191", tis=1.32),
        line("FR0010537423", tis=0.0, typ="TISF"),
        line("FR0011253624", name="R-co Valor", tis=12.65),
        line("FR0011253625", name="R-co Valor", tis=0, typ="TISF"),
    ]


class ValidateTests(unittest.TestCase):
    def test_valid_file_has_no_error(self):
        self.assertEqual(rules.validate_tis_rows(good_rows(), TARGET, RULES), [])

    def test_empty_file_is_refused(self):
        errors = rules.validate_tis_rows([], TARGET, RULES)
        self.assertTrue(any("Aucune ligne" in e for e in errors))

    def test_stale_dates_are_detected(self):
        rows = good_rows()
        rows[0] = line("FR0010523191", day="20261005")
        errors = rules.validate_tis_rows(rows, TARGET, RULES)
        self.assertTrue(any("20261005" in e for e in errors))

    def test_duplicates_are_detected(self):
        rows = good_rows() + [good_rows()[0]]
        errors = rules.validate_tis_rows(rows, TARGET, RULES)
        self.assertTrue(any("en double" in e for e in errors))

    def test_bad_isin_currency_and_type(self):
        rows = good_rows()
        rows[0] = line("XX", ccy="JPY", typ="ZZZZ")
        errors = " ".join(rules.validate_tis_rows(rows, TARGET, RULES))
        self.assertIn("ISIN invalide", errors)
        self.assertIn("JPY", errors)
        self.assertIn("ZZZZ", errors)

    def test_row_count_bounds(self):
        strict = dict(RULES, min_rows=10)
        errors = rules.validate_tis_rows(good_rows(), TARGET, strict)
        self.assertTrue(any("hors de la plage" in e for e in errors))

    def test_non_numeric_tis_is_detected(self):
        rows = good_rows()
        rows[1] = line("FR0010537423", tis="n/a", typ="TISF")
        errors = rules.validate_tis_rows(rows, TARGET, RULES)
        self.assertTrue(any("non numérique" in e for e in errors))


class ExtractTests(unittest.TestCase):
    HEADER = ("ISIN", "FundName / Description", "Currency", "Update Date",
              "NAV Estimated", "NAV Final", "MONTANT TIS", "TYPE FISCALITE", None, None)

    def test_extracts_rows_under_header(self):
        sheet = [(None,) * 10, ("Rothschild & Co Asset Management",) + (None,) * 9,
                 (None,) * 10, self.HEADER] + [line() + (None, None)]
        rows = rules.extract_tis_rows(sheet)
        self.assertEqual(len(rows), 1)
        self.assertEqual(len(rows[0]), 8)

    def test_wrong_header_is_refused(self):
        bad = [("ISIN", "Nom", "Devise")]
        with self.assertRaises(rules.ValidationError):
            rules.extract_tis_rows(bad)

    def test_missing_header_is_refused(self):
        with self.assertRaises(rules.ValidationError):
            rules.extract_tis_rows([("autre",)])


class ValorTests(unittest.TestCase):
    def test_only_exact_name_is_selected(self):
        rows = good_rows() + [line("FR0099999999", name="R-co Valor Balanced")]
        selected = rules.select_valor_rows(rows, "R-co Valor")
        self.assertEqual(len(selected), 2)
        self.assertTrue(all(r[1] == "R-co Valor" for r in selected))

    def test_date_becomes_number_and_blank_nav_becomes_none(self):
        selected = rules.select_valor_rows(good_rows(), "R-co Valor")
        self.assertEqual(selected[0][3], 20261016)
        self.assertIsNone(selected[0][4])

    def test_count_date_accepts_text_and_numbers(self):
        self.assertEqual(rules.count_date(["20261016", 20261016, 20261005, None], TARGET), 2)


class CalendarTests(unittest.TestCase):
    def test_weekend_is_not_business_day(self):
        self.assertFalse(rules.is_business_day(date(2026, 10, 17), []))  # samedi
        self.assertTrue(rules.is_business_day(date(2026, 10, 16), []))   # vendredi

    def test_skip_dates_are_respected(self):
        self.assertFalse(rules.is_business_day(date(2026, 10, 16), ["20261016"]))

    def test_date_key(self):
        self.assertEqual(rules.date_key(date(2026, 1, 5)), "20260105")


SAMPLES = os.environ.get("TIS_SAMPLE_DIR")


@unittest.skipUnless(SAMPLES, "TIS_SAMPLE_DIR non défini")
class SampleFileTests(unittest.TestCase):
    """Vérifie la validation sur les vrais fichiers d'exemple."""

    def _rows(self, name):
        wb = openpyxl.load_workbook(Path(SAMPLES) / name, read_only=True, data_only=True)
        try:
            return rules.extract_tis_rows(wb["TIS BE"].iter_rows(values_only=True))
        finally:
            wb.close()

    def test_am_file_of_16_oct_is_valid(self):
        cfg = {"min_rows": 150, "max_rows": 400, "currencies": ["EUR", "USD", "CHF", "GBP"],
               "types": ["TIBR", "TISF"], "allow_duplicates": False}
        rows = self._rows("7a2aa24f-TISBE_Rothschild__Co_AM_File.xlsx")
        self.assertEqual(rules.validate_tis_rows(rows, "20261016", cfg), [])

    def test_stale_prod_file_is_refused_for_16_oct(self):
        cfg = {"min_rows": 150, "max_rows": 400, "currencies": ["EUR", "USD", "CHF", "GBP"],
               "types": ["TIBR", "TISF"], "allow_duplicates": False}
        rows = self._rows("733d60dc-TISBE_PROD.xlsx")
        errors = " | ".join(rules.validate_tis_rows(rows, "20261016", cfg))
        self.assertIn("20261005", errors)
        self.assertIn("en double", errors)

    def test_valor_history_has_30_rows_per_day_for_exact_name(self):
        wb = openpyxl.load_workbook(Path(SAMPLES) / "d2fb8b79-TIS_BE_R-co_Valor.xlsx",
                                    read_only=True, data_only=True)
        try:
            values = [r for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True)]
        finally:
            wb.close()
        self.assertEqual(rules.count_date((r[3] for r in values), "20261016"), 30)


if __name__ == "__main__":
    unittest.main()
