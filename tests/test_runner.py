# -*- coding: utf-8 -*-
"""Orchestration avec des agents simulés : reprises, chat neuf, reprise après arrêt, disjoncteur."""
from __future__ import annotations

import pytest

from pia_bench import report as R
from pia_bench import scoring as S
from pia_bench.catalogue import load_catalogue
from pia_bench.cli import main
from pia_bench.runner import CampaignStopped, Runner, bind_catalogue, open_campaign, plan_units
from pia_bench.transports.base import AnswerTimeout, Conversation, Reply, Transport
from conftest import make_workbook


class ScriptedTransport(Transport):
    """Répond juste ; `script` permet d'injecter une réponse/erreur pour (n° de chat, n° de tour)."""
    name = "scripted"

    def __init__(self, catalogue, script=None, default=None):
        self.by_text = {q.question: q for q in catalogue.questions}
        self.script = script or {}
        self.default = default
        self.chats = []
        self.recovered = 0

    def new_conversation(self):
        t = self

        class Chat(Conversation):
            def __init__(self):
                self.n = len(t.chats) + 1
                self.asked = []
                self.url = f"chat://{self.n}"
                t.chats.append(self)

            def ask(self, question, timeout, stable):
                self.asked.append(question)
                action = t.script.get((self.n, len(self.asked)))
                if isinstance(action, BaseException):
                    raise action
                if isinstance(action, str):
                    return Reply(action, 1.0)
                if t.default:
                    return Reply(t.default, 1.0)
                q = t.by_text[question]
                text = f"Résultat {q.attendu}" if q.notable else "Résultat Aucune information disponible."
                return Reply(text, 1.0)

        return Chat()

    def recover(self, hard=False):
        self.recovered += 1


def setup(cfg):
    cat = load_catalogue(cfg.path(cfg["catalogue"]["path"]))
    folder, store = open_campaign(cfg, "test")
    bind_catalogue(store, cat)
    return cat, folder, store


def test_one_fresh_chat_per_conversation(cfg):
    cat, folder, store = setup(cfg)
    t = ScriptedTransport(cat)
    stats = Runner(cfg, store, t, folder, sleep=lambda s: None).run(plan_units(cat, store))
    assert stats["complete"] == 4 and stats["failed"] == 0
    assert len(t.chats) == 4
    assert all(len(c.asked) == 5 for c in t.chats)
    assert [len({q.split(']')[0] for q in c.asked}) for c in t.chats] == [1, 1, 1, 1]  # jamais deux conversations par chat
    R.score_campaign(cfg, cat, store)
    rows = R.collect(cat, store)
    assert {r.verdict for r in rows} == {S.CORRECT, S.ABSTENTION}


def test_technical_error_replays_whole_conversation_in_new_chat(cfg):
    cat, folder, store = setup(cfg)
    t = ScriptedTransport(cat, script={(2, 3): "Oops! Something happened. Can you try again?",
                                       (3, 1): AnswerTimeout("trop long")})
    Runner(cfg, store, t, folder, sleep=lambda s: None).run(plan_units(cat, store))
    # C001 : chat 1 · C002 : chat 2 (Oops au tour 3), chat 3 (timeout au tour 1), chat 4 OK · C003, C004
    assert len(t.chats) == 6
    assert [len(c.asked) for c in t.chats[:4]] == [5, 3, 1, 5]
    attempts = [(a["cid"], a["attempt_no"], a["status"]) for a in store.attempts()]
    assert ("C002", 1, "FAILED") in attempts and ("C002", 2, "FAILED") in attempts and ("C002", 3, "COMPLETE") in attempts
    final = {(t_["qid"]): t_["attempt_no"] for t_ in store.final_turns()}
    assert final["Q0008"] == 3  # le tour 3 de C002 vient de la tentative complète
    assert t.recovered == 2


def test_resume_skips_completed_and_replays_interrupted(cfg):
    cat, folder, store = setup(cfg)
    t = ScriptedTransport(cat, script={(3, 2): KeyboardInterrupt()})
    with pytest.raises(KeyboardInterrupt):
        Runner(cfg, store, t, folder, sleep=lambda s: None).run(plan_units(cat, store))
    assert store.completed_units() == {("C001", 1), ("C002", 1)}
    units = plan_units(cat, store)
    assert [u.conv.cid for u in units] == ["C003", "C004"]
    t2 = ScriptedTransport(cat)
    Runner(cfg, store, t2, folder, sleep=lambda s: None).run(units)
    assert len(store.completed_units()) == 4
    assert t2.chats[0].asked[0].startswith("[C003]")


def test_catalogue_change_is_refused(cfg, tmp_path):
    cat, folder, store = setup(cfg)
    other = load_catalogue(make_workbook(tmp_path / "autre.xlsx", n_conv=2))
    with pytest.raises(CampaignStopped):
        bind_catalogue(store, other)


def test_breaker_stops_cleanly_when_agent_is_down(cfg):
    cfg["run"].update(max_attempts=1, breaker_threshold=2, breaker_max_trips=0)
    cat, folder, store = setup(cfg)
    t = ScriptedTransport(cat, default="Oops! Something happened. Can you try again?")
    with pytest.raises(CampaignStopped):
        Runner(cfg, store, t, folder, sleep=lambda s: None).run(plan_units(cat, store))
    assert len(t.chats) == 2


def test_degraded_session_is_detected_and_replayed(cfg):
    cfg["run"].update(degraded_window=8, degraded_ratio=0.8)
    cat, folder, store = setup(cfg)
    t = ScriptedTransport(cat, default="Résultat Aucune information disponible pour cette demande.")
    with pytest.raises(CampaignStopped):
        Runner(cfg, store, t, folder, sleep=lambda s: None).run(plan_units(cat, store))
    statuses = [a["status"] for a in store.attempts()]
    assert "SUSPECT" in statuses
    assert t.recovered >= 1


def test_repeat_and_stability(cfg):
    cat, folder, store = setup(cfg)
    t = ScriptedTransport(cat)
    Runner(cfg, store, t, folder, sleep=lambda s: None).run(plan_units(cat, store, repeat=2))
    assert len(t.chats) == 8
    R.score_campaign(cfg, cat, store)
    st = R.summary(R.collect(cat, store), store)["stability"]
    assert st["repetitions"] == 2 and st["pass_hat_k"] == 1.0


def test_cli_end_to_end_with_fake_agent(cfg, tmp_path, monkeypatch):
    toml = tmp_path / "pia_bench.toml"
    toml.write_text(f'''
[catalogue]
path = "{cfg['catalogue']['path']}"
[campaign]
runs_dir = "runs"
[scoring]
equivalences = "{cfg['scoring']['equivalences']}"
[run]
think_time = 0
between_conversations = 0
backoff_base = 0
breaker_cooldown = 0
degraded_window = 0
[fake]
tech_error_rate = 0.1
''', encoding="utf-8")
    assert main(["--config", str(toml), "run", "--transport", "fake", "--campaign", "demo"]) == 0
    folder = tmp_path / "runs" / "demo"
    assert (folder / "resultats.xlsx").is_file() and (folder / "rapport.html").is_file()
    # relecture humaine : on corrige un verdict dans l'Excel puis on le réimporte
    from openpyxl import load_workbook
    wb = load_workbook(folder / "resultats.xlsx")
    ws = wb["RESULTATS"]
    header = [c.value for c in ws[1]]
    ws.cell(row=2, column=header.index("Verdict humain (à remplir)") + 1, value=S.INCORRECT)
    wb.save(folder / "resultats.xlsx")
    assert main(["--config", str(toml), "import-review", "--campaign", "demo"]) == 0
    assert main(["--config", str(toml), "status", "--campaign", "demo"]) == 0
    assert main(["--config", str(toml), "run", "--transport", "fake", "--campaign", "demo"]) == 0  # rien à refaire
