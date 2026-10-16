# -*- coding: utf-8 -*-
"""Notation d'une campagne, indicateurs et rapports (Excel + HTML).

Priorité des verdicts : humain > juge LLM > règles.
Intervalles de confiance : bootstrap sur les CONVERSATIONS (les 5 tours d'une
conversation ne sont pas indépendants), comme dans le rapport du 29/09.
"""
from __future__ import annotations

import html
import json
import random
import statistics
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from . import scoring as S
from .catalogue import Catalogue
from .store import Store

PRIORITY = ("human", "llm", "rules")
SCORED = (S.CORRECT, S.NO_INFO, S.INCORRECT, S.OFF_TOPIC)   # dénominateur de l'exactitude (questions NOTABLE notées)


@dataclass
class Row:
    qid: str
    cid: str
    tour: int
    repetition: int
    attempt_id: int | None
    attempt_no: int | None
    attempt_status: str
    question: str
    produit: str
    champ: str
    theme: str
    groupe: str
    notable: bool
    attendu: str
    answer: str | None
    status: str
    latency: float | None
    finished_at: str | None
    verdict: str = ""
    source: str = ""
    reason: str = ""
    flags: list[str] = field(default_factory=list)


def competitors_by_field(catalogue: Catalogue) -> dict[str, set[str]]:
    """Valeurs connues par champ, et par famille de champs (« Management Expertise Level 0/1/2 »)."""
    by_field: dict[str, set[str]] = defaultdict(set)
    for q in catalogue.questions:
        if q.notable and q.attendu and len(q.attendu) <= 60:
            by_field[q.champ].add(q.attendu)
    family: dict[str, set[str]] = defaultdict(set)
    for champ, values in by_field.items():
        family[" ".join(champ.split()[:3])] |= values
    return {champ: by_field[champ] | family[" ".join(champ.split()[:3])] for champ in by_field}


def score_campaign(cfg, catalogue: Catalogue, store: Store) -> int:
    eq = S.Equivalences.load(cfg.path(cfg["scoring"]["equivalences"]))
    comp = competitors_by_field(catalogue)
    qs = catalogue.by_qid()
    out = []
    for t in store.final_turns():
        q = qs.get(t["qid"])
        if q is None:
            continue
        v = S.score(q.attendu, q.notable, t["answer"], q.champ, q.produit, eq, comp.get(q.champ, ()), t["status"])
        out.append((q.qid, t["repetition"], "rules", t["attempt_id"], v.verdict, v.reason, {"flags": v.flags}))
    store.save_verdicts(out)
    return len(out)


def collect(catalogue: Catalogue, store: Store, include_missing: bool = True) -> list[Row]:
    verdicts = store.verdicts()
    turns = {(t["qid"], t["repetition"]): t for t in store.final_turns()}
    reps = sorted({k[1] for k in turns} or {1})
    rows = []
    for q in catalogue.questions:
        for rep in reps:
            t = turns.get((q.qid, rep))
            if t is None and not include_missing:
                continue
            r = Row(q.qid, q.cid, q.tour, rep, t["attempt_id"] if t else None, t["attempt_no"] if t else None,
                    t["attempt_status"] if t else "", q.question, q.produit, q.champ, q.theme, q.groupe, q.notable,
                    q.attendu, t["answer"] if t else None, t["status"] if t else "NON ENVOYÉ",
                    t["latency"] if t else None, t["finished_at"] if t else None)
            if t is None:
                r.verdict, r.source = S.NOT_CAPTURED, "rules"
            for src in PRIORITY:
                v = verdicts.get((q.qid, rep, src))
                if v is not None and t is not None and v["attempt_id"] in (None, t["attempt_id"]):
                    r.verdict, r.source, r.reason = v["verdict"], src, v["reason"] or ""
                    r.flags = json.loads(v["detail"] or "{}").get("flags", [])
                    break
            rows.append(r)
    return rows


# --- indicateurs ---------------------------------------------------------------
def rates(rows: list[Row]) -> dict:
    c = Counter(r.verdict for r in rows)
    scored = sum(c[v] for v in SCORED)
    answered = c[S.CORRECT] + c[S.INCORRECT]
    empty = c[S.ABSTENTION] + c[S.INVENTED]
    return {
        "n": len(rows), "counts": dict(c), "scored": scored,
        "accuracy": c[S.CORRECT] / scored if scored else None,
        "no_info": c[S.NO_INFO] / scored if scored else None,
        "incorrect": c[S.INCORRECT] / scored if scored else None,
        "precision": c[S.CORRECT] / answered if answered else None,
        "abstention": c[S.ABSTENTION] / empty if empty else None,
    }


def bootstrap(rows: list[Row], metric: str = "accuracy", n: int = 2000, seed: int = 7) -> tuple[float, float] | None:
    by_conv: dict[tuple[str, int], list[Row]] = defaultdict(list)
    for r in rows:
        by_conv[(r.cid, r.repetition)].append(r)
    keys = list(by_conv)
    if len(keys) < 5:
        return None
    rng = random.Random(seed)
    values = []
    for _ in range(n):
        sample = [r for _ in keys for r in by_conv[rng.choice(keys)]]
        v = rates(sample)[metric]
        if v is not None:
            values.append(v)
    if not values:
        return None
    values.sort()
    return values[int(0.025 * len(values))], values[int(0.975 * len(values)) - 1]


def latency_stats(rows: list[Row]) -> dict:
    lat = sorted(r.latency for r in rows if r.latency is not None and r.status == "CAPTURED")
    if not lat:
        return {}
    pick = lambda p: lat[min(len(lat) - 1, int(p * len(lat)))]
    return {"n": len(lat), "p50": statistics.median(lat), "p90": pick(0.90), "p95": pick(0.95), "max": lat[-1],
            "mean": statistics.fmean(lat)}


def conversation_scores(rows: list[Row]) -> Counter:
    """Nombre de tours notables corrects par conversation complète (toutes les questions notées)."""
    by_conv: dict[tuple[str, int], list[Row]] = defaultdict(list)
    for r in rows:
        by_conv[(r.cid, r.repetition)].append(r)
    out = Counter()
    for turns in by_conv.values():
        notable = [r for r in turns if r.notable]
        if notable and all(r.verdict in SCORED for r in notable):
            out[f"{sum(r.verdict == S.CORRECT for r in notable)}/{len(notable)}"] += 1
    return out


def grouped(rows: list[Row], key) -> list[tuple[str, dict]]:
    groups: dict[str, list[Row]] = defaultdict(list)
    for r in rows:
        groups[str(key(r))].append(r)
    return sorted(((k, rates(v)) for k, v in groups.items()), key=lambda kv: kv[0])


def stability(rows: list[Row]) -> dict | None:
    by_q: dict[str, list[str]] = defaultdict(list)
    for r in rows:
        if r.notable and r.verdict in SCORED:
            by_q[r.qid].append(r.verdict)
    multi = {k: v for k, v in by_q.items() if len(v) > 1}
    if not multi:
        return None
    k = max(len(v) for v in multi.values())
    full = {q: v for q, v in multi.items() if len(v) == k}
    return {
        "questions": len(full), "repetitions": k,
        "same_verdict": sum(len(set(v)) == 1 for v in full.values()) / len(full),
        "pass_at_k": sum(S.CORRECT in v for v in full.values()) / len(full),     # au moins une fois juste
        "pass_hat_k": sum(all(x == S.CORRECT for x in v) for v in full.values()) / len(full),  # juste à chaque fois
        "unstable": sorted(q for q, v in full.items() if len(set(v)) > 1),
    }


def summary(rows: list[Row], store: Store) -> dict:
    sent = [r for r in rows if r.verdict != S.NOT_CAPTURED]
    attempts = store.attempts()
    reasons = Counter((a["error"] or "").split(":")[0] for a in attempts if a["status"] == "FAILED")
    return {
        "questions": len({r.qid for r in rows}),
        "captured": len(sent),
        "overall": rates(sent),
        "ci_accuracy": bootstrap(sent, "accuracy"),
        "ci_precision": bootstrap(sent, "precision"),
        "latency": latency_stats(sent),
        "conversations": conversation_scores(sent),
        "by_tour": grouped(sent, lambda r: r.tour),
        "by_field": grouped(sent, lambda r: r.champ),
        "by_theme": grouped(sent, lambda r: r.theme or "—"),
        "by_group": grouped(sent, lambda r: r.groupe or "—"),
        "attempts": {"total": len(attempts), "status": dict(Counter(a["status"] for a in attempts)),
                     "failure_reasons": dict(reasons),
                     "retried_conversations": len({(a["cid"], a["repetition"]) for a in attempts if a["attempt_no"] > 1})},
        "stability": stability(sent),
        "flags": dict(Counter(f for r in sent for f in r.flags)),
    }


# --- sorties ----------------------------------------------------------------------
def pct(x) -> str:
    return "—" if x is None else f"{100 * x:.1f} %"


def write_xlsx(path: Path, rows: list[Row], summ: dict, store: Store) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    bold = Font(bold=True)
    yellow = PatternFill("solid", fgColor="FFF2CC")
    head = PatternFill("solid", fgColor="DDEBF7")

    def sheet(title, header, data, widths=None, wrap_cols=()):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = bold, head
        for d in data:
            ws.append(list(d))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for col in wrap_cols:
            for cell in ws[get_column_letter(col)][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        return ws

    o = summ["overall"]
    ws = wb.active
    ws.title = "SYNTHESE"
    lines = [
        ("Questions dans le benchmark", summ["questions"]),
        ("Réponses capturées", summ["captured"]),
        ("Réponses notables notées", o["scored"]),
        ("Exactitude", pct(o["accuracy"])),
        ("Exactitude IC 95 % (bootstrap par conversation)", _ci(summ["ci_accuracy"])),
        ("« Pas d’information » à tort", pct(o["no_info"])),
        ("Incorrectes", pct(o["incorrect"])),
        ("Précision quand l’agent répond", pct(o["precision"])),
        ("Précision IC 95 %", _ci(summ["ci_precision"])),
        ("Abstention correcte sur cellule vide", pct(o["abstention"])),
        ("À confirmer (relecture / juge)", o["counts"].get(S.REVIEW, 0) + o["counts"].get(S.INVENTED, 0)),
        ("Échecs techniques restants", o["counts"].get(S.TECH, 0)),
        ("Latence médiane (s)", _f(summ["latency"].get("p50"))),
        ("Latence p95 (s)", _f(summ["latency"].get("p95"))),
        ("Tentatives (chats ouverts)", summ["attempts"]["total"]),
        ("Conversations rejouées au moins une fois", summ["attempts"]["retried_conversations"]),
    ]
    ws.append(["Indicateur", "Valeur"])
    for c in ws[1]:
        c.font, c.fill = bold, head
    for line in lines:
        ws.append(list(line))
    ws.append([])
    ws.append(["Score par conversation (tours notables justes)", "Conversations"])
    for k, v in sorted(summ["conversations"].items(), reverse=True):
        ws.append([k, v])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 52, 24

    res = sheet("RESULTATS",
                ["ID question", "Conversation", "Tour", "Répétition", "Groupe", "Thème", "Produit", "Champ",
                 "Statut GT", "Question posée", "Valeur attendue (base)", "Réponse de l’agent", "Verdict",
                 "Source du verdict", "Motif", "Signaux", "Latence (s)", "Tentative", "Capture (UTC)",
                 "Verdict humain (à remplir)", "Commentaire relecteur"],
                [(r.qid, r.cid, r.tour, r.repetition, r.groupe, r.theme, r.produit, r.champ,
                  "NOTABLE" if r.notable else "EXPLORATOIRE", r.question, r.attendu, r.answer, r.verdict, r.source,
                  r.reason, ", ".join(r.flags), r.latency, r.attempt_no, r.finished_at,
                  r.verdict if r.source == "human" else None, None) for r in rows],
                [9, 8, 5, 6, 12, 16, 28, 26, 11, 50, 30, 70, 18, 10, 30, 14, 9, 8, 20, 18, 30], (10, 11, 12))
    for row in res.iter_rows(min_row=2, min_col=20, max_col=21):
        for c in row:
            c.fill = yellow

    rate_head = ["Réponses", "Notées", "Correctes", "Pas d’info à tort", "Incorrectes", "% correct", "% pas d’info",
                 "À confirmer"]

    def rate_rows(groups):
        return [(k, v["n"], v["scored"], v["counts"].get(S.CORRECT, 0), v["counts"].get(S.NO_INFO, 0),
                 v["counts"].get(S.INCORRECT, 0), v["accuracy"], v["no_info"], v["counts"].get(S.REVIEW, 0))
                for k, v in groups]

    for title, key, label in (("PAR_TOUR", "by_tour", "Tour"), ("PAR_CHAMP", "by_field", "Champ"),
                              ("PAR_THEME", "by_theme", "Thème"), ("PAR_GROUPE", "by_group", "Groupe")):
        ws = sheet(title, [label] + rate_head, rate_rows(summ[key]), [40, 10, 10, 10, 14, 11, 11, 12, 12])
        for row in ws.iter_rows(min_row=2, min_col=7, max_col=8):
            for c in row:
                c.number_format = "0.0%"

    sheet("TENTATIVES", ["Id", "Conversation", "Répétition", "Tentative", "Statut", "Début", "Fin", "Erreur",
                         "Transport", "URL du chat", "Capture d’écran"],
          [(a["id"], a["cid"], a["repetition"], a["attempt_no"], a["status"], a["started"], a["finished"], a["error"],
            a["transport"], a["chat_url"], a["screenshot"]) for a in store.attempts()],
          [6, 10, 8, 8, 10, 22, 22, 60, 10, 40, 24])

    st = summ["stability"]
    if st:
        sheet("STABILITE", ["Indicateur", "Valeur"], [
            ("Questions répétées", st["questions"]), ("Répétitions", st["repetitions"]),
            ("Même verdict à chaque fois", pct(st["same_verdict"])),
            ("pass@k : juste au moins une fois", pct(st["pass_at_k"])),
            ("pass^k : juste à chaque fois", pct(st["pass_hat_k"])),
            ("Questions instables", ", ".join(st["unstable"][:200])),
        ], [40, 80])
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)


def _ci(ci) -> str:
    return "—" if not ci else f"{100 * ci[0]:.1f} % – {100 * ci[1]:.1f} %"


def _f(x) -> str:
    return "—" if x is None else f"{x:.1f}"


def write_html(path: Path, summ: dict, campaign: str) -> None:
    o = summ["overall"]
    e = html.escape

    def kpi(label, value, sub=""):
        return f'<div class="kpi"><div class="v">{e(str(value))}</div><div class="l">{e(label)}</div><div class="s">{e(sub)}</div></div>'

    def bar(row):
        k, v = row
        c = v["counts"]
        total = v["scored"] or 1
        seg = "".join(f'<span class="{cls}" style="width:{100 * c.get(name, 0) / total:.2f}%"></span>'
                      for name, cls in ((S.CORRECT, "ok"), (S.NO_INFO, "ni"), (S.INCORRECT, "ko")))
        return (f"<tr><td>{e(k)}</td><td class=n>{v['scored']}</td><td class=n>{pct(v['accuracy'])}</td>"
                f"<td class=n>{pct(v['no_info'])}</td><td class=n>{c.get(S.INCORRECT, 0)}</td>"
                f"<td class=n>{c.get(S.REVIEW, 0)}</td><td class=bar><div>{seg}</div></td></tr>")

    def table(title, groups, label, sort_by_acc=False):
        groups = sorted(groups, key=lambda kv: (kv[1]["accuracy"] is None, kv[1]["accuracy"] or 0)) if sort_by_acc else groups
        body = "".join(bar(g) for g in groups if g[1]["scored"])
        return (f"<h2>{e(title)}</h2><table><thead><tr><th>{e(label)}</th><th>Notées</th><th>Correct</th>"
                f"<th>Pas d’info</th><th>Faux</th><th>À confirmer</th><th></th></tr></thead><tbody>{body}</tbody></table>")

    conv = "".join(f"<tr><td>{e(k)}</td><td class=n>{v}</td></tr>"
                   for k, v in sorted(summ["conversations"].items(), reverse=True))
    att = summ["attempts"]
    fails = "".join(f"<li>{e(k or 'inconnu')} : {v}</li>" for k, v in att["failure_reasons"].items()) or "<li>aucun</li>"
    st = summ["stability"]
    stab = (f"<h2>Stabilité ({st['repetitions']} répétitions)</h2><p>Même verdict à chaque fois : {pct(st['same_verdict'])}"
            f" · juste au moins une fois (pass@k) : {pct(st['pass_at_k'])} · juste à chaque fois (pass^k) : "
            f"{pct(st['pass_hat_k'])}</p>") if st else ""
    lat = summ["latency"]
    doc = f"""<!doctype html><html lang="fr"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>Benchmark Product Info</title>
<style>
:root{{--bg:#fff;--fg:#1d2433;--mut:#5b6475;--card:#f4f6fa;--line:#e2e6ee;--ok:#2e7d4f;--ni:#c99a2e;--ko:#c0392b}}
@media (prefers-color-scheme:dark){{:root{{--bg:#14171d;--fg:#e8ebf1;--mut:#9aa3b2;--card:#1d222b;--line:#2c323d}}}}
body{{background:var(--bg);color:var(--fg);font:15px/1.5 system-ui,Segoe UI,sans-serif;margin:0 auto;max-width:1100px;padding:16px}}
h1{{font-size:22px}}h2{{font-size:17px;margin-top:28px}}p,li{{color:var(--mut)}}
.kpis{{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:10px}}
.kpi{{background:var(--card);border-radius:10px;padding:12px}}.kpi .v{{font-size:24px;font-weight:650}}
.kpi .l{{font-size:13px}}.kpi .s{{font-size:12px;color:var(--mut)}}
table{{border-collapse:collapse;width:100%;font-size:13px}}td,th{{border-bottom:1px solid var(--line);padding:5px 6px;text-align:left}}
.n{{text-align:right;font-variant-numeric:tabular-nums}}.bar div{{display:flex;height:10px;width:180px;background:var(--line);border-radius:5px;overflow:hidden}}
.ok{{background:var(--ok)}}.ni{{background:var(--ni)}}.ko{{background:var(--ko)}}
.legend span{{display:inline-block;width:10px;height:10px;border-radius:2px;margin:0 4px 0 12px}}
.wrap{{overflow-x:auto}}
</style></head><body>
<h1>Product Info Assistant — campagne {e(campaign)}</h1>
<p>{summ['captured']} réponses capturées sur {summ['questions']} questions · un chat neuf par conversation ·
verdicts : règles, puis juge LLM et relecture humaine quand ils existent.</p>
<div class="kpis">
{kpi("Exactitude", pct(o['accuracy']), "IC 95 % " + _ci(summ['ci_accuracy']))}
{kpi("« Aucune info » à tort", pct(o['no_info']))}
{kpi("Précision quand il répond", pct(o['precision']), "IC 95 % " + _ci(summ['ci_precision']))}
{kpi("Abstention correcte", pct(o['abstention']), "sur cellules vides")}
{kpi("Latence médiane", _f(lat.get('p50')) + " s", "p95 " + _f(lat.get('p95')) + " s")}
{kpi("À confirmer", o['counts'].get(S.REVIEW, 0) + o['counts'].get(S.INVENTED, 0), "relecture ou `judge`")}
</div>
<p class="legend"><span class="ok"></span>correct<span class="ni"></span>aucune info à tort<span class="ko"></span>faux</p>
<div class="wrap">{table("Par position dans la conversation (tour)", summ['by_tour'], "Tour")}</div>
<div class="wrap">{table("Par champ (du plus faible au plus fort)", summ['by_field'], "Champ", True)}</div>
<div class="wrap">{table("Par thème", summ['by_theme'], "Thème", True)}</div>
<h2>Score par conversation</h2><table><tbody>{conv}</tbody></table>
{stab}
<h2>Fiabilité de la collecte</h2>
<p>{att['total']} chats ouverts · statuts : {e(json.dumps(att['status'], ensure_ascii=False))} ·
{att['retried_conversations']} conversation(s) rejouée(s) dans un chat neuf.</p><ul>{fails}</ul>
</body></html>"""
    path.write_text(doc, encoding="utf-8")


def import_review(xlsx: Path, store: Store) -> int:
    """Relit la colonne « Verdict humain (à remplir) » de RESULTATS et l'enregistre (source humaine)."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    rows = wb["RESULTATS"].iter_rows(values_only=True)
    header = list(next(rows))
    i_q, i_rep = header.index("ID question"), header.index("Répétition")
    i_v, i_c = header.index("Verdict humain (à remplir)"), header.index("Commentaire relecteur")
    turns = {(t["qid"], t["repetition"]): t["attempt_id"] for t in store.final_turns()}
    out = []
    for r in rows:
        verdict = (r[i_v] or "").strip() if isinstance(r[i_v], str) else r[i_v]
        if not verdict:
            continue
        if verdict not in S.ALL_VERDICTS:
            raise ValueError(f"{r[i_q]} : verdict humain inconnu {verdict!r} (valeurs : {S.ALL_VERDICTS})")
        rep = int(r[i_rep] or 1)
        out.append((r[i_q], rep, "human", turns.get((r[i_q], rep)), verdict, r[i_c] or "relecture", {}))
    store.save_verdicts(out)
    wb.close()
    return len(out)
