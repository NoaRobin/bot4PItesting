# -*- coding: utf-8 -*-
"""Ligne de commande : python -m pia_bench <commande> [options]

  check          contrôle le classeur (orphelines recollées, conversations, tours)
  setup          ouvre Edge, vous ouvrez l'agent une fois : l'URL est mémorisée
  run            lance / reprend une campagne (un chat neuf par conversation)
  score          recalcule les pré-verdicts par règles
  judge          fait trancher les « À confirmer » par un LLM (optionnel)
  report         produit resultats.xlsx et rapport.html
  import-review  relit les verdicts humains saisis dans resultats.xlsx
  calibrate      compare les règles à des verdicts humains existants
  status         avancement de la campagne
"""
from __future__ import annotations

import argparse
import logging
import sys
from collections import Counter
from pathlib import Path

from . import report as R
from . import scoring as S
from .catalogue import load_catalogue
from .config import load_config
from .runner import CampaignStopped, Runner, bind_catalogue, open_campaign, plan_units
from .transports import make_transport

log = logging.getLogger("pia_bench")


def setup_logging(folder: Path | None = None, verbose: bool = False) -> None:
    log.setLevel(logging.DEBUG if verbose else logging.INFO)
    log.handlers.clear()
    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(logging.Formatter("%(asctime)s %(message)s", "%H:%M:%S"))
    log.addHandler(console)
    if folder:
        folder.mkdir(parents=True, exist_ok=True)
        fh = logging.FileHandler(folder / "run.log", encoding="utf-8")
        fh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
        log.addHandler(fh)


def _catalogue(cfg):
    c = cfg["catalogue"]
    return load_catalogue(cfg.path(c["path"]), c["sheet"], int(c["expected_questions"] or 0))


def cmd_check(cfg, a) -> int:
    cat = _catalogue(cfg)
    convs = cat.conversations
    print(f"Classeur : {cat.path.name}  (SHA-256 {cat.digest[:12]}…)")
    print(f"{len(cat.questions)} questions · {len(convs)} conversations · "
          f"{sum(q.notable for q in cat.questions)} NOTABLE · {sum(not q.notable for q in cat.questions)} EXPLORATOIRE")
    print(f"Lignes orphelines recollées : {cat.merged_orphans}")
    print("Tours par conversation :", dict(Counter(len(c.turns) for c in convs)))
    for w in cat.warnings[:20]:
        print("  ⚠", w)
    first = convs[0].turns[0]
    print(f"Première question : {first.qid} {first.cid} T{first.tour} : {first.question[:100]}")
    return 0


def cmd_setup(cfg, a) -> int:
    from .transports.edge_ui import EdgeUITransport
    t = EdgeUITransport(cfg, None, interactive=True)
    from playwright.sync_api import sync_playwright
    t.pw = sync_playwright().start()
    try:
        t._launch()
        if t.agent_url:
            t.page.goto(t.agent_url)
        t.setup()
        print("Test d'ouverture d'un chat neuf…")
        t.new_conversation()
        print("OK : le script sait ouvrir un chat neuf tout seul. Lancez maintenant `python -m pia_bench run`.")
    finally:
        t.close()
    return 0


def cmd_run(cfg, a) -> int:
    if a.transport:
        cfg["run"]["transport"] = a.transport
    for key in ("max_attempts", "turn_timeout", "stable_seconds", "think_time"):
        if getattr(a, key, None) is not None:
            cfg["run"][key] = getattr(a, key)
    cat = _catalogue(cfg)
    folder, store = open_campaign(cfg, a.campaign, a.resume)
    setup_logging(folder, a.verbose)
    try:
        bind_catalogue(store, cat)
        aborted = store.recover_running()
        if aborted:
            log.info("%d tentative(s) interrompue(s) lors du dernier arrêt : conversations rejouées", aborted)
        units = plan_units(cat, store, a.conversations, a.start_from, a.limit, a.repeat)
        log.info("Campagne %s · transport %s · %d conversation(s) à traiter (%d questions)",
                 folder.name, cfg["run"]["transport"], len(units), sum(len(u.conv.turns) for u in units))
        if not units:
            log.info("Rien à faire : tout est capturé. `report` pour les résultats.")
        else:
            transport = make_transport(cfg["run"]["transport"], cfg, folder, cat, interactive=not a.unattended)
            try:
                transport.start()
                stats = Runner(cfg, store, transport, folder).run(units)
                log.info("Fin : %s", stats)
            except CampaignStopped as e:
                log.error("ARRÊT : %s", e)
                return 2
            except KeyboardInterrupt:
                log.warning("Interrompu (Ctrl+C) : tout ce qui est capturé est conservé ; reprise : run --resume")
                return 130
            finally:
                transport.close()
        _report(cfg, cat, store, folder)
        return 0
    finally:
        store.close()


def _report(cfg, cat, store, folder) -> dict:
    R.score_campaign(cfg, cat, store)
    rows = R.collect(cat, store)
    summ = R.summary(rows, store)
    R.write_xlsx(folder / "resultats.xlsx", rows, summ, store)
    R.write_html(folder / "rapport.html", summ, folder.name)
    o = summ["overall"]
    log.info("Exactitude %s (IC 95 %% %s) · pas d’info à tort %s · précision %s · à confirmer %d",
             R.pct(o["accuracy"]), R._ci(summ["ci_accuracy"]), R.pct(o["no_info"]), R.pct(o["precision"]),
             o["counts"].get(S.REVIEW, 0) + o["counts"].get(S.INVENTED, 0))
    log.info("Rapports : %s et %s", folder / "resultats.xlsx", folder / "rapport.html")
    return summ


def _open_existing(cfg, a):
    folder, store = open_campaign(cfg, a.campaign, resume=not a.campaign)
    setup_logging(folder, getattr(a, "verbose", False))
    cat = _catalogue(cfg)
    bind_catalogue(store, cat)
    return cat, folder, store


def cmd_score(cfg, a) -> int:
    cat, folder, store = _open_existing(cfg, a)
    n = R.score_campaign(cfg, cat, store)
    log.info("%d réponses notées par règles", n)
    store.close()
    return 0


def cmd_report(cfg, a) -> int:
    cat, folder, store = _open_existing(cfg, a)
    _report(cfg, cat, store, folder)
    store.close()
    return 0


def cmd_judge(cfg, a) -> int:
    from .judge import judge_rows, needs_judge
    cat, folder, store = _open_existing(cfg, a)
    R.score_campaign(cfg, cat, store)
    rows = [r for r in R.collect(cat, store, include_missing=False)
            if r.source == "rules" and needs_judge(r.verdict, not a.all)]
    if a.limit:
        rows = rows[: a.limit]
    log.info("%d réponse(s) envoyée(s) au juge (%s)", len(rows), cfg["judge"]["model"])
    payload = [{"qid": r.qid, "question": r.question, "champ": r.champ, "produit": r.produit, "attendu": r.attendu,
                "notable": r.notable, "answer": r.answer, "repetition": r.repetition, "attempt_id": r.attempt_id}
               for r in rows]
    results = judge_rows(cfg, payload)
    store.save_verdicts((r["qid"], r["repetition"], "llm", r["attempt_id"], v, j, {}) for r, v, j in results)
    _report(cfg, cat, store, folder)
    store.close()
    return 0


def cmd_import_review(cfg, a) -> int:
    cat, folder, store = _open_existing(cfg, a)
    n = R.import_review(Path(a.xlsx) if a.xlsx else folder / "resultats.xlsx", store)
    log.info("%d verdict(s) humain(s) importé(s)", n)
    _report(cfg, cat, store, folder)
    store.close()
    return 0


def cmd_status(cfg, a) -> int:
    cat, folder, store = _open_existing(cfg, a)
    done = store.completed_units()
    total = len(cat.conversations)
    statuses = Counter(x["status"] for x in store.attempts())
    print(f"Campagne {folder.name} : {len({c for c, _ in done})}/{total} conversations complètes · tentatives {dict(statuses)}")
    store.close()
    return 0


def cmd_calibrate(cfg, a) -> int:
    """Compare les règles aux verdicts humains d'un classeur RESULTATS (format du 29/09)."""
    from openpyxl import load_workbook
    setup_logging()
    wb = load_workbook(a.xlsx, read_only=True, data_only=True)
    rows = list(wb[a.sheet].iter_rows(values_only=True))
    h = [str(x or "") for x in rows[0]]
    col = lambda *names: next(h.index(n) for n in names if n in h)
    i_exp, i_ans = col("Valeur attendue (base)"), col("Réponse de l’agent", "Réponse de l'agent")
    i_v, i_gt, i_ch, i_pr, i_q = col("Verdict"), col("Statut GT"), col("Champ"), col("Produit"), col("ID question")
    data = [r for r in rows[1:] if r[i_q]]
    by_field: dict[str, set[str]] = {}
    for r in data:
        if str(r[i_gt]).upper() != "EXPLORATOIRE" and len(str(r[i_exp])) <= 60:
            by_field.setdefault(r[i_ch], set()).add(str(r[i_exp]))
    family: dict[str, set[str]] = {}
    for champ, values in by_field.items():
        family.setdefault(" ".join(str(champ).split()[:3]), set()).update(values)
    comp = {champ: values | family[" ".join(str(champ).split()[:3])] for champ, values in by_field.items()}
    eq = S.Equivalences.load(cfg.path(cfg["scoring"]["equivalences"]))
    matrix, diffs = Counter(), []
    for r in data:
        notable = str(r[i_gt]).upper() != "EXPLORATOIRE"
        v = S.score(str(r[i_exp]) if notable else "", notable, str(r[i_ans] or ""), r[i_ch], r[i_pr], eq,
                    comp.get(r[i_ch], ()))
        matrix[(r[i_v], v.verdict)] += 1
        if r[i_v] != v.verdict:
            diffs.append((r[i_q], r[i_v], v.verdict, v.reason))
    agree = sum(n for (h_, m), n in matrix.items() if h_ == m)
    false_correct = sum(n for (h_, m), n in matrix.items() if m == S.CORRECT and h_ != S.CORRECT)
    print(f"Accord règles / humain : {agree}/{len(data)} ({100 * agree / len(data):.1f} %) · "
          f"faux « Correct » : {false_correct}")
    print("\nHumain → Règles")
    for (h_, m), n in sorted(matrix.items(), key=lambda kv: -kv[1]):
        print(f"  {n:4d}  {h_:22s} → {m}")
    print("\nÉcarts :")
    for d in diffs:
        print("  ", " | ".join(map(str, d)))
    return 0


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(prog="python -m pia_bench", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", help="fichier TOML (défaut : pia_bench.toml s'il existe)")
    sub = ap.add_subparsers(dest="cmd", required=True)

    sub.add_parser("check", help="contrôler le classeur")
    sub.add_parser("setup", help="mémoriser l'URL de l'agent (une fois)")

    p = sub.add_parser("run", help="lancer ou reprendre une campagne")
    p.add_argument("--campaign", help="nom de la campagne (créée ou reprise)")
    p.add_argument("--resume", action="store_true", help="reprendre la dernière campagne")
    p.add_argument("--conversations", help="sélection, ex. C001-C020,C143")
    p.add_argument("--from", dest="start_from", help="commencer à cette conversation")
    p.add_argument("--limit", type=int, help="nombre max de conversations (pilote)")
    p.add_argument("--repeat", type=int, default=1, help="passages par conversation (stabilité, pass^k)")
    p.add_argument("--transport", choices=["edge", "directline", "fake"])
    p.add_argument("--max-attempts", type=int)
    p.add_argument("--turn-timeout", type=int)
    p.add_argument("--stable-seconds", type=float)
    p.add_argument("--think-time", type=float)
    p.add_argument("--unattended", action="store_true", help="ne jamais attendre l'opérateur (échec si connexion requise)")
    p.add_argument("-v", "--verbose", action="store_true")

    for name in ("score", "report", "status"):
        q = sub.add_parser(name)
        q.add_argument("--campaign")
    q = sub.add_parser("judge", help="juge LLM sur les cas à confirmer")
    q.add_argument("--campaign")
    q.add_argument("--all", action="store_true", help="juger toutes les réponses, pas seulement les incertaines")
    q.add_argument("--limit", type=int)
    q = sub.add_parser("import-review", help="importer les verdicts humains")
    q.add_argument("--campaign")
    q.add_argument("--xlsx")
    q = sub.add_parser("calibrate", help="comparer les règles à des verdicts humains")
    q.add_argument("xlsx")
    q.add_argument("--sheet", default="RESULTATS")

    a = ap.parse_args(argv)
    cfg = load_config(a.config)
    if a.cmd not in ("run",):
        setup_logging()
    handlers = {"check": cmd_check, "setup": cmd_setup, "run": cmd_run, "score": cmd_score, "report": cmd_report,
                "judge": cmd_judge, "import-review": cmd_import_review, "status": cmd_status,
                "calibrate": cmd_calibrate}
    try:
        return handlers[a.cmd](cfg, a)
    except (CampaignStopped, FileNotFoundError, ValueError) as e:
        print("ERREUR :", e)
        return 1
