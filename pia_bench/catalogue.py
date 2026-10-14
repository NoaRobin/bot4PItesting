# -*- coding: utf-8 -*-
"""Lecture du classeur de questions et regroupement par conversation.

Corrige à la lecture les défauts relevés dans AUDIT_BENCHMARK :
- lignes « orphelines » (textes longs éclatés sur plusieurs lignes Excel) :
  recollées à la ligne précédente au lieu d'être ignorées ;
- empreinte SHA-256 du contenu : une campagne ne reprend jamais sur un autre classeur.
"""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Any

REQUIRED = ["ID question", "ID conversation", "Tour", "Question utilisateur"]
KNOWN = {
    "ID question": "qid",
    "ID conversation": "cid",
    "Tour": "tour",
    "Question utilisateur": "question",
    "Produit cible exact": "produit",
    "Champ exact": "champ",
    "Valeur source brute": "attendu",
    "Statut GT": "gt",
}


def clean(value: Any) -> str:
    return re.sub(r"[ \t\xa0]+", " ", str(value if value is not None else "")).strip()


@dataclass
class Question:
    qid: str
    cid: str
    tour: int
    question: str
    produit: str = ""
    champ: str = ""
    attendu: str = ""
    gt: str = ""
    extra: dict[str, str] = field(default_factory=dict)

    @property
    def notable(self) -> bool:
        """Vrai si la base contient une valeur (statut NOTABLE / cellule non vide)."""
        if self.gt:
            return self.gt.upper() != "EXPLORATOIRE"
        return bool(self.attendu)

    @property
    def theme(self) -> str:
        for key in ("Thème", "Theme", "Thème du scénario", "Scénario"):
            if self.extra.get(key):
                return self.extra[key]
        return ""

    @property
    def groupe(self) -> str:
        return self.extra.get("Groupe", "")


@dataclass
class Conversation:
    cid: str
    turns: list[Question]


@dataclass
class Catalogue:
    path: Path
    questions: list[Question]
    digest: str
    merged_orphans: int = 0
    warnings: list[str] = field(default_factory=list)

    @property
    def conversations(self) -> list[Conversation]:
        order: list[str] = []
        groups: dict[str, list[Question]] = {}
        for q in self.questions:
            if q.cid not in groups:
                order.append(q.cid)
                groups[q.cid] = []
            groups[q.cid].append(q)
        return [Conversation(cid, sorted(groups[cid], key=lambda q: q.tour)) for cid in order]

    def by_qid(self) -> dict[str, Question]:
        return {q.qid: q for q in self.questions}


def load_catalogue(path: Path, sheet: str = "QUESTIONS_1500", expected: int = 0) -> Catalogue:
    from openpyxl import load_workbook

    if not path.is_file():
        raise FileNotFoundError(f"Classeur introuvable : {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Onglet {sheet!r} absent. Onglets : {wb.sheetnames}")
        rows = wb[sheet].iter_rows(values_only=True)
        headers = [clean(h) for h in next(rows)]
        for key in REQUIRED:
            if headers.count(key) != 1:
                raise ValueError(f"Colonne {key!r} absente ou dupliquée")
        questions: list[Question] = []
        merged = 0
        for values in rows:
            raw = {h: v for h, v in zip(headers, values) if h}
            qid = clean(raw.get("ID question"))
            if not qid:
                # Ligne orpheline : suite d'un texte long de la ligne précédente.
                parts = {h: str(v).strip() for h, v in raw.items() if v is not None and str(v).strip()}
                if parts and questions:
                    _append_orphan(questions[-1], parts)
                    merged += 1
                continue
            q = Question(qid=qid, cid="", tour=0, question="")
            for header, value in raw.items():
                attr = KNOWN.get(header)
                text = str(value).strip() if value is not None else ""
                if attr == "tour":
                    q.tour = int(float(text or 0))
                elif attr:
                    setattr(q, attr, clean(text) if attr != "attendu" else text)
                elif text:
                    q.extra[header] = text
            if q.question:
                questions.append(q)
    finally:
        wb.close()

    warnings = _validate(questions, expected)
    payload = json.dumps([asdict(q) for q in questions], ensure_ascii=False, sort_keys=True)
    digest = hashlib.sha256(payload.encode("utf-8")).hexdigest()
    return Catalogue(path=path, questions=questions, digest=digest, merged_orphans=merged, warnings=warnings)


def _append_orphan(q: Question, parts: dict[str, str]) -> None:
    for header, text in parts.items():
        attr = KNOWN.get(header)
        if attr in ("attendu", "question"):
            setattr(q, attr, (getattr(q, attr) + "\n" + text).strip())
        elif attr is None:
            q.extra[header] = (q.extra.get(header, "") + "\n" + text).strip()


def _validate(questions: list[Question], expected: int) -> list[str]:
    if not questions:
        raise ValueError("Aucune question dans le classeur")
    ids = [q.qid for q in questions]
    dup = sorted({x for x in ids if ids.count(x) > 1})
    if dup:
        raise ValueError(f"ID question dupliqués : {dup[:10]}")
    if expected and len(questions) != expected:
        raise ValueError(f"{len(questions)} questions lues, {expected} attendues")
    warnings = []
    seen_cid: dict[str, list[int]] = {}
    for q in questions:
        seen_cid.setdefault(q.cid, []).append(q.tour)
    for cid, tours in seen_cid.items():
        if sorted(tours) != list(range(1, len(tours) + 1)):
            warnings.append(f"{cid} : tours {sorted(tours)} non contigus")
    return warnings
